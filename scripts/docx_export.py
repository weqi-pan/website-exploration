from __future__ import annotations

from pathlib import Path

try:  # python-docx is optional; the runtime skips DOCX export when it is absent.
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError:  # pragma: no cover - depends on the environment
    Document = None


def _set_cn_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def export_function_list_docx(
    pages: list[dict],
    system_name: str,
    out_path: Path,
    assets_base: Path | None = None,
) -> Path | None:
    if Document is None:
        return None
    document = Document()
    _set_cn_font(document)

    document.add_heading("附录一：功能测试结果", level=1)
    document.add_heading("（一）功能测试明细", level=2)
    document.add_heading(f"1. {system_name}", level=3)
    document.add_heading("（1）功能清单", level=3)

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("功能模块", "功能项", "项目说明", "备注")):
        cell.text = text

    for page in pages:
        row = table.add_row().cells
        row[0].text = page.get("moduleName") or ""
        row[1].text = page.get("name") or ""
        row[2].text = page.get("functionDescription") or ""
        row[3].text = page.get("remarks") or ""

    document.add_heading("（2）系统截图", level=3)
    shot_added = False
    for index, page in enumerate(pages, 1):
        rel = page.get("screenshot")
        if not rel:
            continue
        image_path = (assets_base / rel) if assets_base else Path(rel)
        if not image_path.exists():
            continue
        heading = document.add_paragraph(f"{index}）{page.get('name', '')}")
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        try:
            document.add_picture(str(image_path), width=Inches(5.9))
            shot_added = True
        except Exception:
            continue
    if not shot_added:
        document.add_paragraph("本次探索未生成可嵌入的页面截图。")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    return out_path
